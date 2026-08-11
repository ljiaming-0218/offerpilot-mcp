import re
from pathlib import Path

from backend.errors import AppError
from backend.services import llm_service


MAX_RESUME_SIZE = 5 * 1024 * 1024

SKILL_KEYWORDS = [
    "Python", "Java", "JavaScript", "TypeScript", "C++", "Go", "FastAPI",
    "Spring Boot", "MySQL", "MongoDB", "Redis", "SQL", "RAG", "Agent",
    "MCP", "Text2SQL", "LangChain", "LlamaIndex", "Prompt", "Embedding",
    "Vector DB", "Milvus", "FAISS", "Docker", "Linux", "Git", "Vue",
    "React", "HTML", "CSS", "PyTorch", "TensorFlow", "OpenAI",
]

ROLE_KEYWORDS = [
    "AI应用开发", "大模型应用开发", "后端开发", "Python开发", "Java开发",
    "算法实习生", "数据开发", "RAG工程师", "Agent开发", "软件开发",
]

CITY_KEYWORDS = [
    "北京", "上海", "深圳", "广州", "杭州", "成都", "南京", "苏州",
    "武汉", "西安", "重庆", "长沙", "天津",
]


def parse_resume(filename: str, content: bytes) -> dict:
    if not content:
        raise AppError(
            error="EMPTY_RESUME_FILE",
            message="上传的简历文件为空",
            status_code=400,
        )

    if len(content) > MAX_RESUME_SIZE:
        raise AppError(
            error="RESUME_FILE_TOO_LARGE",
            message="简历文件过大，当前限制为 5MB",
            status_code=400,
        )

    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf":
        text = _extract_pdf_text(content)
    elif suffix == ".docx":
        text = _extract_docx_text(content)
    else:
        raise AppError(
            error="UNSUPPORTED_RESUME_TYPE",
            message="仅支持 PDF 或 DOCX 简历",
            status_code=400,
            filename=filename,
        )

    text = _normalize_text(text)
    if len(text) < 20:
        raise AppError(
            error="RESUME_TEXT_TOO_SHORT",
            message="简历文本过短，可能是扫描件或无法解析的文件",
            status_code=400,
        )

    fallback_profile = _parse_resume_by_rules(text)
    warnings = []

    try:
        profile = _parse_resume_by_llm(text)
        parse_method = "llm"
    except AppError as exc:
        profile = fallback_profile
        parse_method = "rule_fallback"
        warnings.append(f"{exc.error}: {exc.message}")

    return {
        "profile": profile,
        "text_preview": text[:800],
        "source": {
            "filename": filename,
            "file_type": suffix.lstrip("."),
            "text_length": len(text),
            "parse_method": parse_method,
            "warnings": warnings,
        },
    }


def _parse_resume_by_rules(text: str) -> dict:
    return {
        "user_id": None,
        "name": _extract_name(text),
        "education": _extract_education(text),
        "target_roles": _extract_keywords(text, ROLE_KEYWORDS),
        "target_cities": _extract_keywords(text, CITY_KEYWORDS),
        "preferred_work_type": _extract_work_type(text),
        "skills": _extract_keywords(text, SKILL_KEYWORDS),
        "projects": _extract_projects(text),
        "weaknesses": [],
    }


def _parse_resume_by_llm(text: str) -> dict:
    system_prompt = """
你是 OfferPilot 的简历画像抽取器。你只能根据简历原文抽取信息，不能编造。
输出必须是严格 JSON，不要 Markdown，不要解释。
字段要求：
{
  "user_id": null,
  "name": string|null,
  "education": string|null,
  "target_roles": string[],
  "target_cities": string[],
  "preferred_work_type": string|null,
  "skills": string[],
  "projects": string[],
  "weaknesses": []
}
规则：
1. 简历没有明确写出的字段返回 null 或 []。
2. weaknesses 永远返回 []，不要推测候选人短板。
3. skills 提取技术技能、工具、框架、数据库、AI/RAG/Agent 相关能力。
4. projects 每条保留为一句完整项目摘要，包含项目名和核心工作。
5. target_roles 和 target_cities 只提取简历中明确出现的求职意向。
6. education 尽量包含学历、学校、专业；缺失部分不要补。
""".strip()
    user_prompt = f"简历原文：\n{text[:12000]}"
    parsed = llm_service.chat_json(system_prompt=system_prompt, user_prompt=user_prompt)
    return _normalize_profile(parsed)


def _normalize_profile(profile: dict) -> dict:
    return {
        "user_id": None,
        "name": _clean_optional_text(profile.get("name")),
        "education": _clean_optional_text(profile.get("education")),
        "target_roles": _clean_list(profile.get("target_roles")),
        "target_cities": _clean_list(profile.get("target_cities")),
        "preferred_work_type": _clean_optional_text(profile.get("preferred_work_type")),
        "skills": _clean_list(profile.get("skills")),
        "projects": _clean_list(profile.get("projects")),
        "weaknesses": [],
    }


def _clean_optional_text(value) -> str | None:
    if not isinstance(value, str):
        return None
    value = value.strip()
    return value or None


def _clean_list(value) -> list[str]:
    if not isinstance(value, list):
        return []
    result = []
    seen = set()
    for item in value:
        if not isinstance(item, str):
            continue
        cleaned = item.strip()
        key = cleaned.lower()
        if cleaned and key not in seen:
            result.append(cleaned)
            seen.add(key)
    return result


def _extract_pdf_text(content: bytes) -> str:
    try:
        import pymupdf
    except ImportError as exc:
        raise AppError(
            error="RESUME_PDF_PARSER_NOT_INSTALLED",
            message="缺少 PDF 解析依赖，请先安装 PyMuPDF",
            status_code=500,
        ) from exc

    with pymupdf.open(stream=content, filetype="pdf") as document:
        return "\n".join(page.get_text("text") for page in document)


def _extract_docx_text(content: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise AppError(
            error="RESUME_DOCX_PARSER_NOT_INSTALLED",
            message="缺少 DOCX 解析依赖，请先安装 python-docx",
            status_code=500,
        ) from exc

    from io import BytesIO

    document = Document(BytesIO(content))
    paragraphs = [paragraph.text for paragraph in document.paragraphs]
    table_cells = [
        cell.text
        for table in document.tables
        for row in table.rows
        for cell in row.cells
    ]
    return "\n".join([*paragraphs, *table_cells])


def _normalize_text(text: str) -> str:
    lines = [line.strip() for line in text.replace("\r", "\n").split("\n")]
    return "\n".join(line for line in lines if line)


def _extract_name(text: str) -> str | None:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    for line in lines[:8]:
        if len(line) <= 12 and not re.search(r"简历|求职|电话|邮箱|@", line, re.I):
            return line
    return None


def _extract_education(text: str) -> str | None:
    degree_patterns = ["博士", "硕士", "研究生", "本科", "学士", "大专"]
    major_pattern = r"[\u4e00-\u9fa5A-Za-z]+(?:科学与技术|工程|技术|信息|数学|智能|软件|计算机|自动化|电子)"
    degree = next((item for item in degree_patterns if item in text), None)
    major_match = re.search(major_pattern, text)
    if degree and major_match:
        return f"{degree}，{major_match.group(0)}"
    return degree or (major_match.group(0) if major_match else None)


def _extract_work_type(text: str) -> str | None:
    if "实习" in text or "实习生" in text:
        return "实习"
    if "校招" in text or "应届" in text:
        return "校招"
    if "全职" in text:
        return "全职"
    return None


def _extract_keywords(text: str, keywords: list[str]) -> list[str]:
    lower_text = text.lower()
    result = []
    for keyword in keywords:
        if keyword.lower() in lower_text and keyword not in result:
            result.append(keyword)
    return result


def _extract_projects(text: str) -> list[str]:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    projects = []
    capture = False

    for line in lines:
        if re.search(r"项目|Project", line, re.I):
            capture = True
            if len(line) > 6:
                projects.append(line)
            continue

        if capture and len(line) <= 12 and re.search(r"教育|校园|技能|证书|获奖|工作经历|实习经历", line):
            capture = False

        if capture and 8 <= len(line) <= 160:
            projects.append(line)

        if len(projects) >= 5:
            break

    return _dedupe(projects)


def _dedupe(values: list[str]) -> list[str]:
    result = []
    seen = set()
    for value in values:
        key = value.lower()
        if key not in seen:
            result.append(value)
            seen.add(key)
    return result
